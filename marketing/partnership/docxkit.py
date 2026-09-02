#!/usr/bin/env python3
"""Shared Word layout kit for the Lomonec partnership documents.

Both the outward-facing proposal and the internal target list are built from
these, so they stay visually identical: same palette, same ruled section
headings, same lavender cards, same gutter arithmetic.
"""
import pathlib
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Emu

HERE = pathlib.Path(__file__).resolve().parent
A = HERE / "assets"

INK    = RGBColor(0x2B, 0x1F, 0x5E)   # deep indigo — headings, dark cards
ACCENT = RGBColor(0x4A, 0x32, 0xDC)   # bright indigo — rules, eyebrows, numerals
BODY   = RGBColor(0x3A, 0x3A, 0x46)
MUTED  = RGBColor(0x6B, 0x70, 0x83)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
LILAC  = RGBColor(0xA7, 0x98, 0xFF)
CARD   = "EFEDFB"                     # lavender card fill
CARD2  = "F6F5FD"
DARK   = "2B1F5E"
GOLD   = "FFF7E3"
MINT   = "EAF6F0"
ROSE   = "FDEEEE"          # internal-only / caution callouts

PAGE_W = Cm(21.0)
MARGIN = Cm(1.7)
CONTENT_W = Cm(21.0 - 1.7 * 2)




# ---------- low-level Word helpers python-docx does not expose ----------
def shade(cell, hexfill):
    el = OxmlElement("w:shd"); el.set(qn("w:val"), "clear"); el.set(qn("w:fill"), hexfill)
    cell._tc.get_or_add_tcPr().append(el)

def cell_margins(cell, top=140, start=160, bottom=140, end=160):
    mar = OxmlElement("w:tcMar")
    for tag, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        e = OxmlElement("w:" + tag); e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa")
        mar.append(e)
    cell._tc.get_or_add_tcPr().append(mar)

def no_borders(table):
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement("w:" + edge); e.set(qn("w:val"), "none"); e.set(qn("w:sz"), "0")
        b.append(e)
    table._tbl.tblPr.append(b)

def cell_border(cell, colour="DDD9F5", sz=6, edges=("top", "left", "bottom", "right")):
    pr = cell._tc.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    for edge in edges:
        e = OxmlElement("w:" + edge)
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(sz))
        e.set(qn("w:space"), "0"); e.set(qn("w:color"), colour)
        b.append(e)
    pr.append(b)

def spacing(run, twentieths):
    """Letterspacing — the offer sheet's eyebrows depend on it."""
    e = OxmlElement("w:spacing"); e.set(qn("w:val"), str(twentieths))
    run._element.get_or_add_rPr().append(e)

def para_rule(p, colour="4A32DC", sz=10):
    pr = p._p.get_or_add_pPr()
    b = OxmlElement("w:pBdr")
    e = OxmlElement("w:bottom")
    e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(sz))
    e.set(qn("w:space"), "4"); e.set(qn("w:color"), colour)
    b.append(e); pr.append(b)

def keep_with_next(p):
    pr = p._p.get_or_add_pPr()
    e = OxmlElement("w:keepNext"); pr.append(e)


# ---------- typographic building blocks ----------
def txt(container, text="", size=10, bold=False, colour=BODY, before=0, after=6,
        align=None, line=1.32, italic=False, font="Arial"):
    p = container.add_paragraph() if hasattr(container, "add_paragraph") else container
    r = p.add_run(text)
    r.font.name = font; r.font.size = Pt(size); r.font.bold = bold
    r.font.italic = italic; r.font.color.rgb = colour
    pf = p.paragraph_format
    pf.space_before = Pt(before); pf.space_after = Pt(after)
    pf.line_spacing = line
    if align is not None:
        p.alignment = align
    return p

def eyebrow(container, text, colour=ACCENT, size=7.5, before=0, after=3):
    p = txt(container, "", size=size, before=before, after=after)
    r = p.runs[0]; r.text = text.upper()
    r.font.bold = True; r.font.color.rgb = colour; r.font.size = Pt(size)
    spacing(r, 30)
    return p

def section_head(doc, number, title):
    p = txt(doc, "", size=7.5, before=11, after=2)
    r = p.runs[0]; r.text = number
    r.font.bold = True; r.font.size = Pt(7.5); r.font.color.rgb = ACCENT
    spacing(r, 30)
    h = txt(doc, title, size=15, bold=True, colour=INK, before=0, after=5, line=1.1)
    para_rule(h)
    keep_with_next(h)
    return h

def bullet(cell, lead, rest, tick="✓"):
    p = cell.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(0); pf.space_after = Pt(5); pf.line_spacing = 1.28
    pf.left_indent = Cm(0.52); pf.first_line_indent = Cm(-0.52)
    r = p.add_run(tick + "  ")
    r.font.name = "Arial"; r.font.size = Pt(9); r.font.bold = True; r.font.color.rgb = ACCENT
    r = p.add_run(lead)
    r.font.name = "Arial"; r.font.size = Pt(9); r.font.bold = True; r.font.color.rgb = INK
    if rest:
        r = p.add_run(rest)
        r.font.name = "Arial"; r.font.size = Pt(9); r.font.color.rgb = BODY
    return p

def card(cell, fill, title, body, title_colour=INK, body_colour=BODY,
         eyebrow_text=None, border="DDD9F5"):
    shade(cell, fill); cell_margins(cell); cell_border(cell, border)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
    first = cell.paragraphs[0]
    first.text = ""
    if eyebrow_text:
        r = first.add_run(eyebrow_text.upper())
        r.font.name = "Arial"; r.font.bold = True; r.font.size = Pt(7)
        r.font.color.rgb = ACCENT if fill != DARK else LILAC
        spacing(r, 30)
        first.paragraph_format.space_after = Pt(3)
        first.paragraph_format.space_before = Pt(0)
        tp = cell.add_paragraph()
    else:
        tp = first
    r = tp.add_run(title)
    r.font.name = "Arial"; r.font.bold = True; r.font.size = Pt(10.5)
    r.font.color.rgb = title_colour
    tp.paragraph_format.space_before = Pt(0); tp.paragraph_format.space_after = Pt(3)
    tp.paragraph_format.line_spacing = 1.18
    if body:
        bp = cell.add_paragraph()
        r = bp.add_run(body)
        r.font.name = "Arial"; r.font.size = Pt(9); r.font.color.rgb = body_colour
        bp.paragraph_format.space_before = Pt(0); bp.paragraph_format.space_after = Pt(0)
        bp.paragraph_format.line_spacing = 1.32
    return cell

def pin_width(table, widths):
    """Pin the table to the full measure, in percent.

    Fixed layout plus dxa widths is the textbook answer, but a reader that
    autofits anyway still shrink-wraps a table whose cells are short — which is
    what happened to the welder-pathway table. Percentages are honoured on both
    paths, so the table fills the measure regardless of how it is laid out. The
    dxa values stay on the grid as the proportion hint.
    """
    total = sum(int(x) for x in widths)
    pr = table._tbl.tblPr
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); pr.append(lay)
    w = OxmlElement("w:tblW")
    w.set(qn("w:w"), "5000"); w.set(qn("w:type"), "pct")
    pr.append(w)
    for row in table.rows:
        for cell, cw in zip(row.cells, widths):
            tcpr = cell._tc.get_or_add_tcPr()
            for old in tcpr.findall(qn("w:tcW")):
                tcpr.remove(old)
            e = OxmlElement("w:tcW")
            e.set(qn("w:w"), str(round(int(cw) / total * 5000)))
            e.set(qn("w:type"), "pct")
            tcpr.insert(0, e)


def grid(doc, rows, cols, widths=None, after=Pt(8)):
    t = doc.add_table(rows=rows, cols=cols)
    t.autofit = False
    no_borders(t)
    w = [Emu(int(x)) for x in (widths or [CONTENT_W / cols] * cols)]
    for i, col in enumerate(t.columns):
        col.width = w[i]                 # w:tblGrid — what fixed layout reads
    for row in t.rows:
        for i, c in enumerate(row.cells):
            c.width = w[i]               # w:tcW — what the autofit path reads
    pin_width(t, w)
    return t


GAP = Cm(0.34)


def card_grid(doc, rows, cols):
    """Cards with real gutters: spacer columns and rows between them, so each
    card reads as its own box the way the offer sheet's do, rather than as one
    ruled block."""
    cw = Cm((CONTENT_W.cm - GAP.cm * (cols - 1)) / cols)
    widths = []
    for i in range(cols):
        widths.append(cw)
        if i < cols - 1:
            widths.append(GAP)
    t = grid(doc, rows * 2 - 1, len(widths), widths)
    for ri, row in enumerate(t.rows):
        if ri % 2:                                  # spacer row
            row.height = GAP
            for c in row.cells:
                cell_margins(c, 0, 0, 0, 0)
                c.paragraphs[0].paragraph_format.space_after = Pt(0)
                c.paragraphs[0].runs and None
                c.paragraphs[0].add_run("").font.size = Pt(1)
    return t, (lambda r, c: t.cell(r * 2, c * 2))

def gutter(doc, pts=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)
    p.add_run("").font.size = Pt(pts)
    return p

def bleed_image(doc, path, width=PAGE_W):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = -MARGIN; pf.right_indent = -MARGIN
    pf.space_before = Pt(0); pf.space_after = Pt(0); pf.line_spacing = 1.0
    p.add_run().add_picture(str(path), width=width)
    return p


# OOXML fixes the order of the children of w:tblPr and w:tcPr. Appending as we
# went produced [tblBorders, tblLayout, tblW] and [shd, tcMar, tcBorders], which
# a reader may silently drop — the welder-pathway table, the one with short
# cells, fell back to autofit and collapsed to two-thirds of the measure. Sort
# both into schema order before saving.
TBLPR_ORDER = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual",
               "tblStyleRowBandSize", "tblStyleColBandSize", "tblW", "jc",
               "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout",
               "tblCellMar", "tblLook", "tblCaption", "tblDescription"]
TCPR_ORDER = ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders",
              "shd", "noWrap", "tcMar", "textDirection", "tcFitText", "vAlign",
              "hideMark"]


def _sort_pr(pr, order):
    if pr is None:
        return
    rank = {qn("w:" + n): i for i, n in enumerate(order)}
    kids = list(pr)
    kids.sort(key=lambda e: rank.get(e.tag, len(order)))
    for e in kids:
        pr.append(e)


def normalise(doc):
    for t in doc.tables:
        _sort_pr(t._tbl.tblPr, TBLPR_ORDER)
        for row in t.rows:
            for c in row.cells:
                _sort_pr(c._tc.tcPr, TCPR_ORDER)




def new_document():
    """A4, no top margin — the masthead image bleeds to the page edge."""
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Arial"; st.font.size = Pt(10); st.font.color.rgb = BODY
    st.paragraph_format.space_after = Pt(6)
    s = doc.sections[0]
    s.page_width, s.page_height = PAGE_W, Cm(29.7)
    s.top_margin = Cm(0); s.bottom_margin = Cm(1.1)
    s.left_margin = MARGIN; s.right_margin = MARGIN
    return doc


def callout(doc, fill, border, label, label_colour, body_runs, size=9.5):
    """A tinted full-width box: small caps label, then one or more runs."""
    t = grid(doc, 1, 1, [CONTENT_W])
    c = t.cell(0, 0)
    shade(c, fill); cell_margins(c, 140, 170, 140, 170); cell_border(c, border)
    p = c.paragraphs[0]; p.text = ""
    r = p.add_run(label.upper())
    r.font.name = "Arial"; r.font.bold = True; r.font.size = Pt(7.5)
    r.font.color.rgb = label_colour
    spacing(r, 30); p.paragraph_format.space_after = Pt(4)
    p2 = c.add_paragraph()
    for text, bold in body_runs:
        r = p2.add_run(text)
        r.font.name = "Arial"; r.font.size = Pt(size); r.font.bold = bold
        r.font.color.rgb = INK if bold else BODY
    p2.paragraph_format.space_before = Pt(0); p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.38
    return c
