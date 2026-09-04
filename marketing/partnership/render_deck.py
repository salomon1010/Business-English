#!/usr/bin/env python3
"""Render the .pptx to PNGs so the slides can actually be looked at.

Nothing in this environment can display a .pptx — PowerPoint's AppleScript
reports a successful export and writes no files, and qlmanage hangs on the
format. So this walks the saved presentation and draws what is really in it:
every picture, rounded rectangle and text run at its stored position, size and
colour. It is not PowerPoint's typesetter, so line breaks may differ by a word,
but it is read from the file rather than from intent — which is what makes it
worth looking at.
"""
import pathlib, sys, sys
from pptx import Presentation
from pptx.util import Emu
from PIL import Image, ImageDraw, ImageFont

HERE = pathlib.Path(__file__).resolve().parent
DECK = sys.argv[1] if len(sys.argv) > 1 else "Petrocertif-Pilot-Presentation.pptx"
OUT = HERE / "assets" / ("deck-preview-value" if "Value" in DECK else "deck-preview")
OUT.mkdir(parents=True, exist_ok=True)
SCALE = 2.0                                    # px per point
FONTS = {(False, False): "/System/Library/Fonts/Supplemental/Arial.ttf",
         (True, False): "/System/Library/Fonts/Supplemental/Arial Bold.ttf",
         (False, True): "/System/Library/Fonts/Supplemental/Arial Italic.ttf",
         (True, True): "/System/Library/Fonts/Supplemental/Arial Bold Italic.ttf"}
_c = {}


def font(pt, bold, italic):
    k = (round(pt * SCALE), bool(bold), bool(italic))
    if k not in _c:
        _c[k] = ImageFont.truetype(FONTS[(bool(bold), bool(italic))], k[0])
    return _c[k]


def rgb(c, default=(40, 40, 50)):
    try:
        v = c.rgb
        return (v[0], v[1], v[2])
    except Exception:
        return default


def wrap(draw, runs, width_px):
    """Greedy wrap across a paragraph's runs, keeping each run's styling.
    A run may carry hard newlines; those force a break rather than being
    measured, which Pillow refuses to do."""
    lines, cur, cur_w = [], [], 0.0
    for text, f, col in runs:
        for bi, block in enumerate(text.split("\n")):
            if bi:                                  # hard break inside the run
                lines.append(cur); cur, cur_w = [], 0.0
            for word in block.split(" "):
                if not word:
                    continue
                piece = word + " "
                w = draw.textlength(piece, font=f)
                if cur and cur_w + w > width_px:
                    lines.append(cur); cur, cur_w = [], 0.0
                cur.append((piece, f, col)); cur_w += w
    if cur:
        lines.append(cur)
    return lines or [[]]


prs = Presentation(HERE / DECK)
SW = int(Emu(prs.slide_width).pt * SCALE)
SH = int(Emu(prs.slide_height).pt * SCALE)

for idx, slide in enumerate(prs.slides, start=1):
    img = Image.new("RGB", (SW, SH), (255, 255, 255))
    d = ImageDraw.Draw(img, "RGBA")
    for sh in slide.shapes:
        x = Emu(sh.left).pt * SCALE
        y = Emu(sh.top).pt * SCALE
        w = Emu(sh.width).pt * SCALE
        h = Emu(sh.height).pt * SCALE
        if sh.shape_type == 13:                                   # picture
            blob = sh.image.blob
            import io
            pic = Image.open(io.BytesIO(blob)).convert("RGBA")
            pic = pic.resize((max(1, int(w)), max(1, int(h))), Image.LANCZOS)
            img.paste(pic, (int(x), int(y)), pic)
            continue
        if sh.has_text_frame and sh.shape_type is not None and not sh.text_frame.text.strip():
            try:                                                  # a drawn card
                fill = rgb(sh.fill.fore_color, (246, 245, 253))
                line = rgb(sh.line.color, (221, 217, 245))
                d.rounded_rectangle([x, y, x + w, y + h], radius=10 * SCALE,
                                    fill=fill, outline=line, width=max(1, int(SCALE)))
            except Exception:
                pass
            continue
        if not sh.has_text_frame:
            continue
        try:
            fill = rgb(sh.fill.fore_color, None)
            if fill:
                d.rounded_rectangle([x, y, x + w, y + h], radius=10 * SCALE, fill=fill)
        except Exception:
            pass
        tf = sh.text_frame
        cy = y
        for p in tf.paragraphs:
            runs = [(r.text,
                     font(r.font.size.pt if r.font.size else 12, r.font.bold, r.font.italic),
                     rgb(r.font.color, (40, 40, 50)))
                    for r in p.runs if r.text]
            if not runs:
                cy += 6 * SCALE
                continue
            size = max((r.font.size.pt if r.font.size else 12) for r in p.runs if r.text)
            ls = p.line_spacing if isinstance(p.line_spacing, float) else 1.2
            for ln in wrap(d, runs, w):
                lw = sum(d.textlength(t, font=f) for t, f, _ in ln)
                cx = x + (w - lw) / 2 if str(p.alignment) == "CENTER (2)" else x
                for t, f, col in ln:
                    d.text((cx, cy), t, font=f, fill=col)
                    cx += d.textlength(t, font=f)
                cy += size * ls * SCALE
            cy += (p.space_after.pt if p.space_after else 0) * SCALE
    img.save(OUT / ("slide-%02d.png" % idx))

# one contact sheet so the whole deck can be taken in at once
n = len(prs.slides._sldIdLst)
cols = 3
rows = (n + cols - 1) // cols
tw, th = SW // 4, SH // 4
sheet = Image.new("RGB", (cols * tw + (cols + 1) * 12, rows * th + (rows + 1) * 12),
                  (232, 230, 242))
for i in range(n):
    t = Image.open(OUT / ("slide-%02d.png" % (i + 1))).resize((tw, th), Image.LANCZOS)
    r, c = divmod(i, cols)
    sheet.paste(t, (12 + c * (tw + 12), 12 + r * (th + 12)))
sheet.save(OUT / "contact-sheet.png")
print("rendered %d slides + contact sheet -> %s" % (n, OUT.relative_to(HERE)))
