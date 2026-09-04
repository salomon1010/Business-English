#!/usr/bin/env python3
"""Measure every text box in the deck against the space it was given.

PowerPoint cannot be scripted in this environment — it reports a successful
export and writes nothing — and QuickLook hangs on .pptx, so there is no way to
render the slides here. This does the next best thing: it wraps every run at the
box's real width using Arial metrics at the real point size, adds the paragraph
spacing, and flags any box whose text needs more height than it was given.
That catches the failure that actually matters, which is text spilling onto the
element below it.
"""
import pathlib, sys
from pptx import Presentation
from pptx.util import Emu, Pt
from PIL import ImageFont

HERE = pathlib.Path(__file__).resolve().parent
FONTS = {(False, False): "/System/Library/Fonts/Supplemental/Arial.ttf",
         (True, False): "/System/Library/Fonts/Supplemental/Arial Bold.ttf",
         (False, True): "/System/Library/Fonts/Supplemental/Arial Italic.ttf",
         (True, True): "/System/Library/Fonts/Supplemental/Arial Bold Italic.ttf"}
PX = 4
_c = {}


def font(pt, bold, italic):
    k = (round(pt * PX), bool(bold), bool(italic))
    if k not in _c:
        _c[k] = ImageFont.truetype(FONTS[(bool(bold), bool(italic))], k[0])
    return _c[k]


def lines_for(text, pt, bold, italic, width_pt):
    if not text.strip():
        return 1
    f = font(pt, bold, italic)
    limit = width_pt * PX
    n, line = 1, ""
    for w in text.split():
        trial = (line + " " + w).strip()
        if f.getlength(trial) <= limit:
            line = trial
        else:
            n += 1; line = w
    return n


def needed_height_pt(tf, width_pt):
    total = 0.0
    for p in tf.paragraphs:
        runs = p.runs
        if not runs:
            total += 6
            continue
        size = max((r.font.size.pt if r.font.size else 12) for r in runs)
        bold = any(r.font.bold for r in runs)
        italic = any(r.font.italic for r in runs)
        text = "".join(r.text for r in runs)
        ls = p.line_spacing if isinstance(p.line_spacing, float) else 1.2
        n = lines_for(text, size, bold, italic, width_pt)
        total += n * size * ls
        total += p.space_after.pt if p.space_after else 0
    return total


prs = Presentation(HERE / "Petrocertif-Pilot-Presentation.pptx")
SW, SH = Emu(prs.slide_width).pt, Emu(prs.slide_height).pt
print("slide %.0f x %.0f pt · %d slides\n" % (SW, SH, len(prs.slides._sldIdLst)))

problems = 0
for i, s in enumerate(prs.slides, start=1):
    notes = s.notes_slide.notes_text_frame.text.strip() if s.has_notes_slide else ""
    flags = []
    pics = 0
    for sh in s.shapes:
        if sh.shape_type == 13:
            pics += 1
            if Emu(sh.left).pt < -1 or Emu(sh.top).pt < -1 \
               or Emu(sh.left + sh.width).pt > SW + 1 \
               or Emu(sh.top + sh.height).pt > SH + 1:
                flags.append("picture out of bounds")
            continue
        if not sh.has_text_frame:
            continue
        tf = sh.text_frame
        if not tf.text.strip():
            continue
        w = Emu(sh.width).pt - Emu(tf.margin_left).pt - Emu(tf.margin_right).pt
        need = needed_height_pt(tf, w)
        have = Emu(sh.height).pt
        bottom = Emu(sh.top).pt + need
        if need > have + 2:
            flags.append("overflows box by %.0fpt: %r" % (need - have, tf.text[:44]))
        if bottom > SH - 4:
            flags.append("runs off the slide: %r" % tf.text[:44])
    # boxes that fit their own frame can still land on top of each other; the
    # summary line on the learner-activity slide sat on the footer
    boxes = []
    for sh in s.shapes:
        if sh.shape_type == 13 or not sh.has_text_frame or not sh.text_frame.text.strip():
            continue
        tf = sh.text_frame
        wd = Emu(sh.width).pt - Emu(tf.margin_left).pt - Emu(tf.margin_right).pt
        boxes.append((Emu(sh.left).pt, Emu(sh.top).pt, wd,
                      needed_height_pt(tf, wd), tf.text[:30]))
    for a in range(len(boxes)):
        for b in range(a + 1, len(boxes)):
            ax, ay, aw, ah, at = boxes[a]
            bx, by, bw, bh, bt = boxes[b]
            if ax < bx + bw - 2 and bx < ax + aw - 2 and \
               ay < by + bh - 2 and by < ay + ah - 2:
                flags.append("text collides: %r over %r" % (at, bt))
    if not notes:
        flags.append("NO SPEAKER NOTES")
    status = "ok " if not flags else "!! "
    print("%s slide %-2d  %d pictures, %d words of notes" %
          (status, i, pics, len(notes.split())))
    for f in flags:
        problems += 1
        print("        - " + f)

print("\n%s" % ("all slides fit, every slide has notes" if not problems
                else "%d issue(s) to fix" % problems))
sys.exit(0)
